"""ส่งออกข้อมูลเป็น Excel (.xlsx) และ PDF (ภาษาไทย)"""
import io
import os
import tempfile
import urllib.request
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
HEADERS = ["งวด", "วันที่", "รายการ", "บัตร/ประเภท", "ผู้จ่าย", "จำนวนเงิน",
           "ผ่อน งวดที่", "จากทั้งหมด", "ประจำ", "หมายเหตุ"]
FONT_URL = "https://raw.githubusercontent.com/google/fonts/main/ofl/sarabun/{}.ttf"
BLUE = "2563EB"


def plabel(p):
    try:
        y, m = p.split("-")
        return f"{MONTHS[int(m) - 1]} {y}"
    except Exception:
        return p


def is_blank(x):
    return (x or "").strip() in ("", "-")


def pivot(rows, key):
    """(งวด, กลุ่ม, {งวด: {กลุ่ม: ยอด}}) กลุ่มที่ว่าง/'-' ไม่นับ (บิลแบ่งจ่ายมีทั้งแถวบัตรและแถวคน)"""
    periods = sorted({r["period"] for r in rows})
    totals = {}
    for r in rows:
        if not is_blank(r[key]):
            totals[r[key]] = totals.get(r[key], 0.0) + r["amount"]
    groups = sorted(totals, key=lambda g: -totals[g])
    table = {p: {g: 0.0 for g in groups} for p in periods}
    for r in rows:
        if not is_blank(r[key]):
            table[r["period"]][r[key]] += r["amount"]
    return periods, groups, table


def inst_text(r):
    if r.get("inst_no") is None:
        return ""
    return f"{r['inst_no']}/{r['inst_total']}" if r.get("inst_total") else f"#{r['inst_no']}"


# ---------------- Excel ----------------
def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _pivot_sheet(wb, name, rows, key):
    periods, groups, table = pivot(rows, key)
    ws = wb.create_sheet(name)
    ws.append(["งวด"] + groups + ["รวม"])
    _style_header(ws, len(groups) + 2)
    col_tot = [0.0] * (len(groups) + 1)
    for p in periods:
        vals = [table[p][g] for g in groups]
        ws.append([plabel(p)] + vals + [sum(vals)])
        for i, v in enumerate(vals + [sum(vals)]):
            col_tot[i] += v
    ws.append(["รวมทั้งหมด"] + col_tot)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 14
    for i in range(2, len(groups) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.append([])
    ws.append(["หมายเหตุ: ยอดตามบัตรและตามคนนับแยกกัน (บิลแบ่งจ่ายมีทั้งแถวยอดรวมบัตรและแถวแบ่งตามคน)"])
    return ws


def build_xlsx(rows, title="RCC"):
    wb = Workbook()
    ws = wb.active
    ws.title = "รายการ"
    ws.append(HEADERS)
    _style_header(ws, len(HEADERS))
    for r in rows:
        ws.append([r["period"], r["date"], r["item"], r["type"] or None, r["name"] or None, r["amount"],
                   r["inst_no"], r["inst_total"], "ใช่" if r["recurring"] else None, r["description"] or None])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "DD/MM/YYYY"
        row[5].number_format = "#,##0.00"
    for i, w in enumerate([10, 12, 36, 16, 10, 14, 11, 11, 8, 32], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 2)}"
    _pivot_sheet(wb, "สรุปตามผู้จ่าย", rows, "name")
    _pivot_sheet(wb, "สรุปตามบัตร", rows, "type")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
def ensure_fonts():
    """คืน (regular, bold) ของฟอนต์ Sarabun: หาในโฟลเดอร์ fonts/ ก่อน ไม่มีก็ดาวน์โหลดจาก GitHub"""
    here = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
    out = []
    for name in ("Sarabun-Regular", "Sarabun-Bold"):
        path = os.path.join(here, name + ".ttf")
        if not os.path.exists(path):
            path = os.path.join(tempfile.gettempdir(), name + ".ttf")
            if not os.path.exists(path):
                urllib.request.urlretrieve(FONT_URL.format(name), path)
        out.append(path)
    return tuple(out)


def build_pdf(rows, title, salary_info=None, fonts=None):
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    reg, bold = fonts or ensure_fonts()
    now = datetime.now(timezone(timedelta(hours=7)))

    class PDF(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("Sarabun", "", 8)
            self.set_text_color(120)
            self.cell(0, 6, f"RCC Tracker · {now:%d/%m/%Y %H:%M} · หน้า {self.page_no()}", align="C")

    pdf = PDF(orientation="P", unit="mm", format="A4")
    pdf.add_font("Sarabun", "", reg)
    pdf.add_font("Sarabun", "B", bold)
    pdf.set_text_shaping(True)                      # จัดวรรณยุกต์/สระภาษาไทยให้ถูกตำแหน่ง
    pdf.set_margins(10, 12, 10)
    pdf.set_auto_page_break(True, 16)
    pdf.add_page()
    head = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(37, 99, 235))
    money = lambda x: f"{x:,.2f}"

    pdf.set_font("Sarabun", "B", 17)
    pdf.cell(0, 9, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Sarabun", "", 9)
    pdf.set_text_color(100)
    pdf.cell(0, 5, f"สร้างเมื่อ {now:%d/%m/%Y %H:%M} · {len(rows)} รายการ", new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)

    def mini_table(title_, headers, data, widths):
        pdf.set_font("Sarabun", "B", 11)
        pdf.cell(0, 7, title_, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Sarabun", "", 9)
        with pdf.table(col_widths=widths, width=sum(widths), align="L",
                       text_align=("LEFT",) + ("RIGHT",) * (len(widths) - 1),
                       line_height=5.5, headings_style=head, borders_layout="HORIZONTAL_LINES") as t:
            r = t.row()
            for h in headers:
                r.cell(h)
            for d in data:
                r = t.row()
                for v in d:
                    r.cell(v)
        pdf.ln(3)

    if salary_info:
        pdf.set_font("Sarabun", "B", 12)
        pdf.cell(0, 7, f"เงินเดือนคงเหลือ {money(salary_info['remain'])} บาท", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Sarabun", "", 9)
        pdf.cell(0, 5, f"ฐานเงินเดือน {salary_info['salary']:,.0f} − ยอด {salary_info['owner']} "
                       f"{money(salary_info['pop'])}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    periods, persons, ptab = pivot(rows, "name")
    _, cards, ctab = pivot(rows, "type")
    ptot = {g: sum(ptab[p][g] for p in periods) for g in persons}
    ctot = {g: sum(ctab[p][g] for p in periods) for g in cards}
    if persons:
        mini_table("สรุปตามผู้จ่าย", ["ผู้จ่าย", "ยอดรวม (บาท)"], [(g, money(ptot[g])) for g in persons] +
                   [("รวมทุกคน", money(sum(ptot.values())))], (70, 40))
    if cards:
        mini_table("สรุปตามบัตร", ["บัตร", "ยอดรวม (บาท)"], [(g, money(ctot[g])) for g in cards] +
                   [("รวมทุกบัตร", money(sum(ctot.values())))], (70, 40))
    if len(periods) > 1:
        mini_table("สรุปรายงวด", ["งวด", "รวมทุกบัตร", "รวมทุกคน"],
                   [(plabel(p), money(sum(ctab[p].values())), money(sum(ptab[p].values()))) for p in periods],
                   (50, 45, 45))

    widths = (21, 68, 30, 20, 15, 36)
    for p in periods:
        prow = [r for r in rows if r["period"] == p]
        pdf.set_font("Sarabun", "B", 12)
        if pdf.get_y() > 250:
            pdf.add_page()
        pdf.cell(0, 8, f"งวด {plabel(p)}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Sarabun", "", 9)
        with pdf.table(col_widths=widths, text_align=("LEFT", "LEFT", "LEFT", "LEFT", "CENTER", "RIGHT"),
                       line_height=5.2, headings_style=head, borders_layout="HORIZONTAL_LINES") as t:
            r_ = t.row()
            for h in ("วันที่", "รายการ", "บัตร", "ผู้จ่าย", "ผ่อน", "จำนวนเงิน"):
                r_.cell(h)
            for r in prow:
                item = r["item"] + (f"\n{r['description']}" if r["description"] else "")
                r_ = t.row()
                for v in (r["date"].strftime("%d/%m/%Y") if r["date"] else "-", item, r["type"] or "-",
                          r["name"] or "-", inst_text(r), money(r["amount"])):
                    r_.cell(v)
        card = sum(r["amount"] for r in prow if not is_blank(r["type"]))
        person = sum(r["amount"] for r in prow if not is_blank(r["name"]))
        pdf.set_font("Sarabun", "B", 9)
        pdf.cell(0, 6, f"รวมทุกบัตร {money(card)} · รวมทุกคน {money(person)}", align="R",
                 new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    return bytes(pdf.output())
